from __future__ import annotations

import urllib.parse
import webbrowser
from pathlib import Path

import barcode
from barcode.writer import ImageWriter
from openpyxl import load_workbook
from PySide6.QtWidgets import QDialog, QFileDialog, QInputDialog, QMessageBox, QPushButton

from services.supabase_service import SupabaseService
from views.crud import CrudView, Field, RecordDialog


PRODUCT_FIELDS = [
    Field("sku", "SKU", required=True), Field("name", "Nombre", required=True),
    Field("product_type", "Tipo", "choice", True, ("terminado", "materia_prima")),
    Field("category_id", "ID de categoría"), Field("description", "Descripción", "longtext"),
    Field("unit", "Unidad", "choice", True, ("unidad", "metro", "kilo", "cono", "rollo", "docena")),
    Field("cost", "Costo", "decimal"), Field("price", "Precio", "decimal"),
    Field("minimum_stock", "Stock mínimo", "decimal"), Field("supplier_id", "ID de proveedor"),
    Field("is_active", "Activo", "bool"),
]


class InventoryView(CrudView):
    def __init__(self, service: SupabaseService):
        super().__init__(service, title="Inventario", table="products", permission="inventory.write",
            select="id,sku,name,product_type,unit,cost,price,minimum_stock,is_active,created_at,product_variants(id,stock,size,color,barcode)",
            columns=[("sku", "SKU"), ("name", "Producto"), ("product_type", "Tipo"), ("unit", "Unidad"),
                     ("price", "Precio"), ("minimum_stock", "Mínimo"), ("is_active", "Activo")], fields=PRODUCT_FIELDS)
        movement = QPushButton("Registrar movimiento", objectName="gold")
        movement.clicked.connect(self._movement)
        self.layout().itemAt(1).layout().insertWidget(2, movement)
        imp = QPushButton("Importar Excel"); imp.clicked.connect(self._import_excel)
        self.layout().itemAt(1).layout().insertWidget(3, imp)
        code = QPushButton("Código de barras"); code.clicked.connect(self._barcode)
        self.layout().itemAt(1).layout().insertWidget(4, code)
        variant = QPushButton("Nueva variante"); variant.clicked.connect(self._variant)
        self.layout().itemAt(1).layout().insertWidget(5, variant)

    def _movement(self) -> None:
        row = self._selected()
        if not row: QMessageBox.information(self, "Seleccione", "Seleccione el producto del movimiento."); return
        variants = row.get("product_variants") or []
        options = [f"{v.get('size')} / {v.get('color')} — {v.get('id')}" for v in variants]
        selected, ok = QInputDialog.getItem(self, "Variante", "Talla / color", options, 0, False)
        if not ok or not selected: return
        variant_id = selected.rsplit(" — ", 1)[-1]
        fields = [
                  Field("movement_type", "Tipo", "choice", True, ("entrada", "salida", "ajuste_positivo", "ajuste_negativo")),
                  Field("quantity", "Cantidad", "decimal", True), Field("reason", "Motivo", required=True),
                  Field("reference", "Referencia")]
        dialog = RecordDialog("Movimiento de inventario", fields, parent=self)
        if dialog.exec():
            data = dialog.data()
            try:
                self.service.rpc("register_stock_movement", {
                    "p_variant_id": variant_id, "p_movement_type": data["movement_type"],
                    "p_quantity": data["quantity"], "p_reason": data["reason"], "p_reference": data["reference"]})
                QMessageBox.information(self, "Movimiento registrado", "El stock y kardex se actualizaron en una transacción.")
                self.refresh()
            except Exception as exc: QMessageBox.critical(self, "Movimiento rechazado", str(exc))

    def _variant(self) -> None:
        row = self._selected()
        if not row: QMessageBox.information(self, "Seleccione", "Seleccione un producto."); return
        dialog = RecordDialog("Nueva talla / color", [Field("size","Talla",required=True),Field("color","Color",required=True),Field("barcode","Código de barras")], parent=self)
        if dialog.exec():
            try:
                payload = dialog.data(); payload["product_id"] = row["id"]
                self.service.client.table("product_variants").insert(payload).execute(); self.refresh()
            except Exception as exc: QMessageBox.critical(self, "No se creó la variante", str(exc))

    def _import_excel(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Importar productos", "", "Excel (*.xlsx)")
        if not path: return
        try:
            ws = load_workbook(path, read_only=True, data_only=True).active
            headers = [str(c.value or "").strip().lower() for c in next(ws.rows)]
            required = {"sku", "name", "product_type", "unit"}
            if not required.issubset(headers): raise ValueError("Columnas requeridas: sku, name, product_type, unit")
            records = []
            for values in ws.iter_rows(values_only=True):
                row = dict(zip(headers, values))
                if row.get("sku"): records.append({k: v for k, v in row.items() if k in {f.key for f in PRODUCT_FIELDS} and v is not None})
            if records: self.service.client.table("products").upsert(records, on_conflict="sku").execute()
            QMessageBox.information(self, "Importación completa", f"Se procesaron {len(records)} productos."); self.refresh()
        except Exception as exc: QMessageBox.critical(self, "No se pudo importar", str(exc))

    def _barcode(self) -> None:
        row = self._selected()
        if not row: QMessageBox.information(self, "Seleccione", "Seleccione un producto."); return
        path, _ = QFileDialog.getSaveFileName(self, "Guardar código", f"{row['sku']}.png", "PNG (*.png)")
        if path:
            try:
                result = barcode.get("code128", row["sku"], writer=ImageWriter()).save(str(Path(path).with_suffix("")))
                QMessageBox.information(self, "Código generado", result)
            except Exception as exc: QMessageBox.critical(self, "Error", str(exc))


class CustomersView(CrudView):
    def __init__(self, service: SupabaseService):
        super().__init__(service, title="Clientes y CRM", table="customers", permission="customers.write",
            columns=[("document_number", "DNI/RUC"), ("name", "Cliente"), ("phone", "Teléfono"),
                     ("email", "Correo"), ("customer_type", "Tipo"), ("is_active", "Activo")],
            fields=[Field("name", "Nombre / razón social", required=True), Field("document_type", "Documento", "choice", True, ("DNI", "RUC", "CE", "Otro")),
                    Field("document_number", "Número"), Field("phone", "Teléfono"), Field("whatsapp", "WhatsApp"),
                    Field("email", "Correo"), Field("address", "Dirección"), Field("customer_type", "Tipo", "choice", True, ("persona", "empresa", "institución")),
                    Field("notes", "Observaciones", "longtext"), Field("is_active", "Activo", "bool")])
        wa = QPushButton("Abrir WhatsApp"); wa.clicked.connect(self._whatsapp)
        self.layout().itemAt(1).layout().insertWidget(2, wa)

    def _whatsapp(self) -> None:
        row = self._selected()
        if not row: QMessageBox.information(self, "Seleccione", "Seleccione un cliente."); return
        phone = "".join(ch for ch in str(row.get("whatsapp") or row.get("phone") or "") if ch.isdigit())
        if len(phone) == 9: phone = "51" + phone
        webbrowser.open(f"https://wa.me/{phone}?text={urllib.parse.quote('Hola, le escribe Cordón de Oro.')}")


class QuotationsView(CrudView):
    def __init__(self, service: SupabaseService):
      super().__init__(service, title="Cotizaciones", table="quotations", permission="quotations.write",
        select="id,number,customer_id,status,issue_date,valid_until,subtotal,discount,tax,total,estimated_delivery,notes,created_at",
        columns=[("number", "Número"), ("customer_id", "Cliente ID"), ("status", "Estado"), ("issue_date", "Fecha"), ("total", "Total"), ("estimated_delivery", "Entrega")],
        fields=[Field("customer_id", "ID del cliente", required=True), Field("status", "Estado", "choice", True, ("borrador", "enviada", "aprobada", "rechazada", "vencida")),
                Field("issue_date", "Fecha", "date", True), Field("valid_until", "Válida hasta", "date"), Field("subtotal", "Subtotal", "decimal"),
                Field("discount", "Descuento", "decimal"), Field("tax", "Impuesto", "decimal"), Field("total", "Total", "decimal"),
                Field("estimated_delivery", "Entrega estimada", "date"), Field("notes", "Notas", "longtext")])
      convert = QPushButton("Convertir en pedido", objectName="gold"); convert.clicked.connect(self._convert)
      self.layout().itemAt(1).layout().insertWidget(2, convert)

    def _convert(self) -> None:
        row = self._selected()
        if not row: QMessageBox.information(self,"Seleccione","Seleccione una cotización aprobada."); return
        try:
            order_id = self.service.rpc("convert_quotation_to_order", {"p_quotation_id":row["id"]})
            QMessageBox.information(self,"Pedido creado",f"Pedido: {order_id}")
        except Exception as exc: QMessageBox.critical(self,"No se pudo convertir",str(exc))


def quotations_view(service: SupabaseService) -> CrudView:
    return QuotationsView(service)


def orders_view(service: SupabaseService) -> CrudView:
    statuses = ("registrado", "diseño_pendiente", "diseño_aprobado", "materiales_preparados", "corte", "confección", "bordado", "estampado", "control_calidad", "listo_entrega", "entregado", "cancelado")
    view = CrudView(service, title="Pedidos", table="orders", permission="orders.write",
        columns=[("number", "Pedido"), ("customer_id", "Cliente ID"), ("status", "Estado"), ("order_date", "Fecha"), ("delivery_date", "Entrega"), ("total", "Total"), ("balance", "Saldo")],
        fields=[Field("customer_id", "ID del cliente", required=True), Field("status", "Estado", "choice", True, statuses),
                Field("order_date", "Fecha", "date", True), Field("delivery_date", "Entrega", "date"), Field("responsible_id", "ID responsable"),
                Field("total", "Total", "decimal"), Field("advance", "Adelanto", "decimal"), Field("balance", "Saldo", "decimal"), Field("notes", "Notas", "longtext")])
    status = QPushButton("Cambiar estado", objectName="gold")
    def change_status():
        row = view._selected()
        if not row: QMessageBox.information(view,"Seleccione","Seleccione un pedido."); return
        new, ok = QInputDialog.getItem(view,"Etapa del pedido","Nuevo estado",statuses,statuses.index(row["status"]) if row.get("status") in statuses else 0,False)
        if ok:
            try: service.rpc("update_order_status",{"p_order_id":row["id"],"p_status":new,"p_notes":None}); view.refresh()
            except Exception as exc: QMessageBox.critical(view,"No se cambió el estado",str(exc))
    status.clicked.connect(change_status); view.layout().itemAt(1).layout().insertWidget(2,status)
    return view


def production_view(service: SupabaseService) -> CrudView:
    return CrudView(service, title="Producción y bordado", table="embroidery_jobs", permission="production.write",
        columns=[("number", "Orden"), ("customer_id", "Cliente ID"), ("garment_type", "Prenda"), ("quantity", "Cantidad"),
                 ("stitch_count", "Puntadas"), ("operator_id", "Operador"), ("status", "Estado"), ("approved_pieces", "Aprobadas"), ("defective_pieces", "Defectuosas")],
        fields=[Field("customer_id", "ID cliente"), Field("order_id", "ID pedido"), Field("design_path", "Diseño / ruta"),
                Field("garment_type", "Tipo de prenda", required=True), Field("quantity", "Cantidad", "integer", True), Field("thread_colors", "Colores de hilo"),
                Field("stitch_count", "Puntadas", "integer"), Field("operator_id", "ID operador"), Field("start_date", "Inicio", "date"),
                Field("end_date", "Finalización", "date"), Field("estimated_minutes", "Minutos estimados", "integer"), Field("actual_minutes", "Minutos reales", "integer"),
                Field("approved_pieces", "Piezas aprobadas", "integer"), Field("defective_pieces", "Piezas defectuosas", "integer"),
                Field("status", "Estado", "choice", True, ("pendiente", "programado", "en_proceso", "pausado", "terminado", "cancelado"))])


def suppliers_view(service: SupabaseService) -> CrudView:
    return CrudView(service, title="Proveedores", table="suppliers", permission="purchases.write",
        columns=[("ruc", "RUC"), ("name", "Proveedor"), ("contact_name", "Contacto"), ("phone", "Teléfono"), ("email", "Correo"), ("is_active", "Activo")],
        fields=[Field("ruc", "RUC"), Field("name", "Razón social", required=True), Field("contact_name", "Contacto"), Field("phone", "Teléfono"),
                Field("email", "Correo"), Field("address", "Dirección"), Field("notes", "Notas", "longtext"), Field("is_active", "Activo", "bool")])


def purchases_view(service: SupabaseService) -> CrudView:
    return CrudView(service, title="Compras", table="purchases", permission="purchases.write",
        columns=[("number", "Compra"), ("supplier_id", "Proveedor ID"), ("status", "Estado"), ("purchase_date", "Fecha"), ("total", "Total"), ("paid", "Pagado")],
        fields=[Field("supplier_id", "ID proveedor", required=True), Field("status", "Estado", "choice", True, ("borrador", "ordenada", "recibida", "cancelada")),
                Field("purchase_date", "Fecha", "date", True), Field("subtotal", "Subtotal", "decimal"), Field("tax", "Impuesto", "decimal"),
                Field("total", "Total", "decimal"), Field("paid", "Pagado", "decimal"), Field("notes", "Notas", "longtext")])

