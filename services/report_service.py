from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from utils.formatters import display


def export_excel(path: str, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B1F35")
    for row in rows:
        ws.append([display(_nested(row, key)) for key, _ in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(40, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(path)


def export_pdf(path: str, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"CORDÓN DE ORO E.I.R.L. — {title}", styles["Title"])]
    story.append(Paragraph(f"Generado: {datetime.now():%d/%m/%Y %H:%M}", styles["Normal"]))
    story.append(Spacer(1, 0.4 * cm))
    data = [[label for _, label in columns]]
    data.extend([[display(_nested(row, key)) for key, _ in columns] for row in rows])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B1F35")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDE2E8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F6F8")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    doc.build(story)


def _nested(row: dict[str, Any], key: str) -> Any:
    value: Any = row
    for part in key.split("."):
        value = value.get(part) if isinstance(value, dict) else None
    return value

